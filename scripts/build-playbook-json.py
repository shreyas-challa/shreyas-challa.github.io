"""One-off tool: parse the 'Red vs Blue - Space RVB' Excel workbook's Playbook
and Services sheets into scripts/playbook-plaintext.json, the plaintext input
to scripts/encrypt-playbook.mjs. Not part of the app build; run manually
whenever the source spreadsheet changes.

Usage: python scripts/build-playbook-json.py "<path to xlsx>"
"""
import json
import re
import sys

import openpyxl

KNOWN_MACHINES = ["Apollo", "Hubble", "Pathfinder", "Columbia", "Odyssey", "Sputnik", "Sat", "Router", "Wazuh"]


def slugify(s):
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


def parse_playbook_sheet(ws):
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2, max_row=ws.max_row)]

    sections = []
    current = None
    notes_lines = []
    in_notes = False

    for vals in rows:
        col0 = vals[0]
        rest_blank = all(v in (None, "") for v in vals[1:])
        if col0 in (None, "") and rest_blank:
            continue
        if col0 == "NOTES, CORRECTIONS AND TIMELINE":
            in_notes = True
            continue
        if in_notes:
            if col0:
                notes_lines.append(col0)
            continue
        if col0 in ("L1", "L2", "L3", "L4", "L5"):
            if current is None:
                continue
            level, check, target, action, command, why, revert = vals[:7]
            current["levels"].append({
                "level": level,
                "action": (action or "").strip() if isinstance(action, str) else action,
                "command": command,
                "why": why,
                "revert": revert,
                "disabled": bool(action and "do not use" in str(action).lower()),
            })
            continue
        if current:
            sections.append(current)
        current = {"header": col0, "levels": []}
    if current:
        sections.append(current)

    out_sections = []
    for s in sections:
        header = s["header"]
        title, rest = header.split("—", 1) if "—" in header else (header, "")
        title = title.strip()
        parts = [p.strip() for p in rest.split("  |  ") if p.strip()]
        checks, targets, description_parts, warning = [], [], [], None
        for p in parts:
            if p.startswith("Checks:") or p.startswith("Check:"):
                checks = [c.strip() for c in p.split(":", 1)[1].split(",")]
            elif p.startswith("Targets:") or p.startswith("Target:"):
                targets = [t.strip() for t in p.split(":", 1)[1].split(",")]
            elif p.startswith("⚠"):
                warning = p.lstrip("⚠").strip()
            else:
                description_parts.append(p)
        haystack = " ".join(targets) + " " + " ".join(checks) + " " + header
        machines = [m for m in KNOWN_MACHINES if m in haystack]
        out_sections.append({
            "id": slugify(title),
            "title": title,
            "checks": checks,
            "targets": targets,
            "machines": machines,
            "description": "  ".join(description_parts).strip(),
            "warning": warning,
            "levels": s["levels"],
        })

    return out_sections, notes_lines


def parse_notes(notes_lines):
    timeline, changed, bugs, traps, services_fixes = [], [], [], [], []
    mode = None
    for line in notes_lines:
        stripped = line.strip()
        if stripped.startswith("TIMELINE"):
            mode = "timeline"; continue
        if stripped.startswith("WHAT CHANGED FROM LAST YEAR"):
            mode = "changed"; continue
        if stripped.startswith("BUGS FIXED"):
            mode = "bugs"; continue
        if stripped.startswith("NEW TRAPS"):
            mode = "traps"; continue
        if stripped.startswith("TWO THINGS TO FIX"):
            mode = "servicesfix"; continue
        if stripped.startswith("22 scored checks"):
            continue

        bucket = {"timeline": timeline, "changed": changed, "bugs": bugs,
                  "traps": traps, "servicesfix": services_fixes}.get(mode)
        if bucket is None:
            continue
        if mode == "timeline":
            bucket.append(stripped)
        elif stripped.startswith("•"):
            bucket.append(stripped.lstrip("•").strip())
        elif bucket:
            bucket[-1] += " " + stripped
    return timeline, changed, bugs, traps, services_fixes


def parse_services_sheet(ws):
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2, max_row=ws.max_row)]
    machines = {}
    for row in rows:
        host = row[0]
        if not host or "scored checks" in str(host):
            continue
        os_, ip = row[1], row[2]
        machines.setdefault(host, {"os": os_, "ip": ip, "services": []})
        machines[host]["services"].append({
            "service": row[3],
            "scored": row[4] == "Yes",
            "detail": row[5],
        })
    return machines


def strip_em_dashes(obj):
    """The site's style rule bans em dashes everywhere, including quoted
    spreadsheet text. Recursively swap '—' for a plain hyphen."""
    if isinstance(obj, str):
        return obj.replace("—", "-")
    if isinstance(obj, list):
        return [strip_em_dashes(v) for v in obj]
    if isinstance(obj, dict):
        return {k: strip_em_dashes(v) for k, v in obj.items()}
    return obj


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/build-playbook-json.py <path to xlsx>")
        sys.exit(1)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    sections, notes_lines = parse_playbook_sheet(wb["Playbook"])
    timeline, changed, bugs, traps, services_fixes = parse_notes(notes_lines)
    machines = parse_services_sheet(wb["Services"])

    playbook = {
        "meta": {
            "name": "Space RVB - Red Team Playbook",
            "scoredChecksTotal": 22,
            "timeline": timeline,
            "whatChangedFromLastYear": changed,
            "bugsFixedFromLastYear": bugs,
            "newTrapsThisYear": traps,
            "servicesTabFixes": services_fixes,
        },
        "machines": machines,
        "sections": sections,
    }
    playbook = strip_em_dashes(playbook)

    out_path = "scripts/playbook-plaintext.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(playbook, f, ensure_ascii=False, indent=1)
    print(f"wrote {out_path}: {len(sections)} sections, {len(machines)} machines")


if __name__ == "__main__":
    main()
